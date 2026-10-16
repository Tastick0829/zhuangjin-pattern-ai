# -*- coding: utf-8 -*-
"""论文匿名化（v2）：段落级重建策略，解决人名跨 run 拆分导致的漏替。

v1 的缺陷：逐 run 替换，若「李朋威」被 Word 拆成 run["李朋"] + run["威"] 就无法命中。
v2 的策略：拼接段落全文 → 整体替换 → 写回首个 run、清空其余 run。
           例外：若该段落原本只有一个 run（常见于无格式差异的短行），直接改该 run 保留格式。
"""
import sys, io, os, zipfile, shutil

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from docx import Document

PAPER_DIR = r"D:\wendang\GitHub\zhuangjin-pattern-ai\paper"

# 长串优先：整行作者信息先处理，避免被短串拆散
REPLACEMENTS = [
    ("南宁市第三中学 陆李 李朋威 雷小玥", "{学校名称省略} {匿名作者一} {匿名作者二} {匿名作者三}"),
    ("南宁市第三中学", "{学校名称省略}"),
    ("李朋威", "{匿名作者二}"),
    ("雷小玥", "{匿名作者三}"),
    ("陆李", "{匿名作者一}"),
]
KEYWORDS = ["陆李", "李朋威", "雷小玥", "南宁市第三中学"]


def replace_in_text(text: str) -> str:
    for old, rep in REPLACEMENTS:
        if old in text:
            text = text.replace(old, rep)
    return text


def fix_paragraph(p) -> int:
    runs = p.runs
    if not runs:
        return 0
    full = "".join(r.text for r in runs)
    if not full:
        return 0
    new = replace_in_text(full)
    if new == full:
        return 0
    if len(runs) == 1:
        runs[0].text = new
    else:
        runs[0].text = new
        for r in runs[1:]:
            r.text = ""
    return 1


def process(path: str) -> int:
    doc = Document(path)
    n = 0
    for p in doc.paragraphs:
        n += fix_paragraph(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    n += fix_paragraph(p)
    for section in doc.sections:
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                n += fix_paragraph(p)
    doc.save(path)
    return n


def check(path: str) -> dict:
    z = zipfile.ZipFile(path)
    xml = b"".join(z.read(n) for n in z.namelist() if n.endswith(".xml")).decode("utf-8", "ignore")
    return {kw: xml.count(kw) for kw in KEYWORDS}


def para2(path: str):
    doc = Document(path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return paras[1] if len(paras) > 1 else "(无第2段)"


if __name__ == "__main__":
    print("=" * 64)
    print("论文匿名化 v2（段落级重建）")
    print("=" * 64)
    for name in ["paper-main.docx", "paper-early.docx"]:
        path = os.path.join(PAPER_DIR, name)
        if not os.path.exists(path):
            print(f"\n跳过（不存在）: {name}")
            continue
        before = check(path)
        n = process(path)
        after = check(path)
        print(f"\n【{name}】")
        print(f"  处理前关键词计数: {before}")
        print(f"  修改段落数      : {n}")
        print(f"  处理后关键词计数: {after}")
        ok = all(v == 0 for v in after.values())
        print(f"  结论: {'✓ 已彻底清除' if ok else '✗ 仍有残留，需进一步处理'}")
        print(f"  现在的第 2 段   : {para2(path)[:110]}")

    # 对所有 xml 部件做一次穷尽扫描（更严格）
    print("\n" + "=" * 64)
    print("穷尽扫描（所有 xml 部件，含页眉页脚/注释/批注）")
    print("=" * 64)
    for name in ["paper-main.docx", "paper-early.docx"]:
        path = os.path.join(PAPER_DIR, name)
        if not os.path.exists(path):
            continue
        z = zipfile.ZipFile(path)
        total = 0
        for part in z.namelist():
            if not part.endswith(".xml"):
                continue
            txt = z.read(part).decode("utf-8", "ignore")
            for kw in KEYWORDS:
                c = txt.count(kw)
                if c:
                    print(f"  ✗ {name} :: {part} 含 '{kw}' × {c}")
                    total += c
        print(f"  {name}: {'✓ 全部 xml 部件干净' if total == 0 else f'✗ 共 {total} 处残留'}")
